"""Export both table forms and the ordered stages. Requires openpyxl 3.1.x."""
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

p=Path(__file__).resolve().parent
d=json.loads((p/'map.json').read_text())
w=Workbook()
s=w.active;s.title='Порядок этапов'
s.append(['Этап','Что делаем','Результат','Условие перехода','Срок команды'])
for stage in d['stages']:
    s.append([stage['title'],stage['action'],stage['exit'],stage['gate'],stage['duration']])
s.append(['Принцип',d['workflow_note']])
s.append(['Лимит работы',d['rules']['time_limit']])
s.append(['Отметки',d['status']])
for col,width in enumerate([32,70,60,58,23],1):s.column_dimensions[get_column_letter(col)].width=width
s.freeze_panes='B2';s.auto_filter.ref='A1:E10'
for row in range(2,s.max_row+1):s.row_dimensions[row].height=100 if row<11 else 60
fills={'yes':'E3F2E8','no':'FAEBEA','partial':'FFF1CE','unknown':'EDF0F4','na':'EDF0F4'}
for name,ids in [('Музеи — этапы',d['sequence_columns']),('Полная таблица',[c['id'] for c in d['columns']])]:
    s=w.create_sheet(name)
    columns=[c for c in d['columns'] if c['id'] in ids]
    s.append(['№','Музей / тема','Тип строки']+[c['title'] for c in columns])
    for m in d['museums']:
        s.append([m['number'],m['title'],m['kind']]+[m['matrix'][c['id']]['label'] for c in columns])
        r=s.max_row
        s.row_dimensions[r].height=44
        for i,c in enumerate(columns,4):
            x=m['matrix'][c['id']];cell=s.cell(r,i)
            cell.fill=PatternFill('solid',fgColor=fills.get(x['state'],'EDF0F4'))
            cell.comment=Comment(x['note']+'\n\nИсточник: '+x['source']+('\n'+x['source_url'] if x.get('source_url') else ''),'Источник отметки')
    s.column_dimensions['A'].width=7;s.column_dimensions['B'].width=47;s.column_dimensions['C'].width=19
    for i in range(4,s.max_column+1):s.column_dimensions[get_column_letter(i)].width=18
    s.freeze_panes='D2';s.auto_filter.ref=s.dimensions
    s.print_title_cols='A:C';s.print_title_rows='1:1'
for s in w:
    s.sheet_view.showGridLines=False
    s.row_dimensions[1].height=48
    for row in s:
        for c in row:
            c.font=Font(name='Calibri',size=11,color='18334B')
            c.alignment=Alignment(vertical='top',wrap_text=True)
            c.border=Border(bottom=Side(style='hair',color='DCE3EB'))
    for c in s[1]:
        c.fill=PatternFill('solid',fgColor='18334B');c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.sheet_properties.pageSetUpPr.fitToPage=True
    s.print_options.horizontalCentered=True
out=p/'museum-sequence.xlsx';w.save(out)
check=load_workbook(out)
assert len(check.sheetnames)==3
assert check['Музеи — этапы'].max_row==55
assert check['Музеи — этапы'].max_column==21
assert check['Полная таблица'].max_column==47
print('Excel: 3 sheets; 54 museum/theme rows; 18 sequence and 44 full status columns.')
